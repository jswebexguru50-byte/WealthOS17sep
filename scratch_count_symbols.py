import openpyxl
wb = openpyxl.load_workbook(r'outputs/Six_Strategies_90_Sessions_2026-10-01.xlsx', data_only=True)
syms = set()
for sheet in wb.sheetnames:
    ws = wb[sheet]
    header_row = None
    headers = []
    for i in range(1, 15):
        row = [str(cell.value or '').strip().lower() for cell in ws[i]]
        if 'symbol' in row:
            header_row = i
            headers = row
            break
    if not header_row:
        continue
    idx = headers.index('symbol')
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        val = str(row[idx] or '').strip().upper()
        if val and val != 'NONE': syms.add(val)
print(len(syms))
