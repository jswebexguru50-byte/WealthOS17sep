from pathlib import Path
import sqlite3
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
MD_PATH = ROOT / "outputs/fundamental_dossiers/pilot_25/Full_Fundamental_Analysis_25.md"
OUT_PATH = ROOT / "outputs/fundamental_dossiers/pilot_25/Validated_Fundamental_Analysis_25_2026-10-04.docx"
DB_PATH = ROOT / "portfolio.db"

db = sqlite3.connect(DB_PATH)
db.row_factory = sqlite3.Row
filter_rows = {
    str(r["symbol"]).upper(): dict(r)
    for r in db.execute("SELECT * FROM strategy_fundamental_filter_results ORDER BY evaluated_at DESC").fetchall()
}
db.close()

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.6)
section.bottom_margin = Inches(0.6)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(9)

symbol_count = 0
current_symbol = None
def business_text(line: str) -> str:
    """Remove machine-facing markers and make the existing evidence readable."""
    line = line.replace("[AVAILABLE]", "")
    line = line.replace("DATA_INSUFFICIENT", "information not available")
    line = line.replace("[MISSING]", "")
    line = line.replace("evidenceState: information not available", "Overall evidence: incomplete")
    replacements = {
        "revenueGrowth": "Revenue growth",
        "operatingProfit": "Operating profit",
        "debtToEquity": "Debt-to-equity",
        "totalBorrowings": "Total borrowings",
        "cfoToPat": "Cash flow from operations / PAT",
        "cfoToOperatingProfit": "Cash flow from operations / operating profit",
        "freeCashFlow": "Free cash flow",
        "fcfYield": "FCF yield",
        "promoterHolding": "Promoter holding",
        "promoterPledge": "Promoter pledge",
        "fiiHolding": "FII holding",
        "diiHolding": "DII holding",
        "fiiTrend": "FII trend",
        "diiTrend": "DII trend",
        "marginTrend": "Margin trend",
        "latestClose": "Latest price",
        "latestOhlcvDate": "Price date",
        "freshnessStatus": "Price freshness",
        "stockMomentumStatus": "Stock momentum",
        "sectorMomentumStatus": "Sector momentum",
        "gapSignalToClose": "Signal-to-price gap",
        "missingFields": "Missing elements",
        "evidenceCompletenessPct": "Evidence completeness",
    }
    for old, new in replacements.items():
        line = line.replace(old, new)
    line = line.replace("; ", ". ")
    line = line.replace("score:", "score ")
    line = line.replace("status:", "status ")
    line = line.replace("enabled: true", "enabled")
    line = line.replace("enabled: false", "not enabled")
    return line

for raw in MD_PATH.read_text(encoding="utf-8").splitlines():
    line = raw.strip()
    if not line or line == "---":
        continue
    if line.startswith("# "):
        p = doc.add_heading(line[2:], 0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith("## "):
        if symbol_count:
            doc.add_page_break()
        symbol_count += 1
        current_symbol = line[3:].strip().upper()
        doc.add_heading(line[3:], 1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], 2)
        if line.lower().startswith("### detailed fundamentals"):
            row = filter_rows.get(current_symbol, {})
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Shading Accent 1"
            table.rows[0].cells[0].text = "Fundamental filter"
            table.rows[0].cells[1].text = "Observed value"
            table.rows[0].cells[2].text = "Result"
            checks = [
                ("Promoter holding", "promoter_pct", "promoter_pass"),
                ("Profitability across last 8 quarters", "profitable_quarter_count", "profitable_last_8_quarters"),
                ("ROCE", "roce_pct", "roce_pass"),
                ("ROE", "roe_pct", "roe_pass"),
                ("Promoter pledge", "pledged_pct", "no_pledge_pass"),
                ("Institutional participation", "institutional_increasing", "institutional_involvement_pass"),
                ("Operating cash-flow discipline", "cash_flow_pass", "cash_flow_pass"),
            ]
            for label, value_key, pass_key in checks:
                value = row.get(value_key)
                passed = row.get(pass_key)
                value_text = "Not available" if value is None else str(value)
                result = "Not available" if passed is None else ("Pass" if bool(passed) else "Review")
                cells = table.add_row().cells
                cells[0].text = label
                cells[1].text = value_text
                cells[2].text = result
    else:
        # Keep the business report readable: source IDs and implementation labels
        # are retained in the spreadsheet/database, not repeated in the narrative.
        if line.startswith("Provider overview evidence"):
            line = "The company overview is supported by the captured provider record."
        elif line.startswith("Document-search evidence"):
            line = "Relevant filings and company documents have been retained for reference."
        elif line.startswith("Corporate-event evidence"):
            line = "Corporate-event information has been captured; review the event details for material developments."
        elif line.startswith("Risk Management"):
            line = "Risk controls and action readiness are shown below. Any unavailable inputs are listed in the final sentence."
        elif line.startswith("actionReadiness:"):
            line = line.replace("actionReadiness:", "Action readiness: ")
            line = business_text(line)
        elif line.startswith("revenueGrowth:"):
            line = business_text(line)
        elif line.startswith("cfoToPat:"):
            line = business_text(line)
        elif line.startswith("status:") and "evidenceIds" in line:
            # QGLP machine evidence IDs are intentionally omitted from the reader-facing document.
            parts = line.split(" evidenceIds:", 1)[0]
            line = "QGLP assessment: " + business_text(parts) + ". Detailed source IDs remain in the workbook and database."
        elif line.startswith("missingDataChecklist:"):
            line = "Unavailable inputs requiring follow-up: " + business_text(line.split("missingDataChecklist:", 1)[1]).replace("group", "domain")
        else:
            line = business_text(line)
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(4)

doc.save(OUT_PATH)
print(f"WROTE {OUT_PATH}")
print(f"PARAGRAPHS {len(doc.paragraphs)}")
print(f"SYMBOL_SECTIONS {symbol_count}")
