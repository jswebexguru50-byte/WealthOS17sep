from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "17365D"
TEAL = "0F6B78"
LIGHT_BLUE = "DCE6F1"
LIGHT_GREY = "F2F4F7"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_rich_paragraph(document: Document, text: str, style=None, shade_fill=None):
    paragraph = document.add_paragraph(style=style)
    if shade_fill:
        p_pr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), shade_fill)
        p_pr.append(shd)
    chunks = re.split(r"(\*\*.*?\*\*)", text)
    for chunk in chunks:
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        else:
            cleaned = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", chunk)
            paragraph.add_run(cleaned)
    return paragraph


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12, TEAL)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    if "Evidence" not in document.styles:
        evidence = document.styles.add_style("Evidence", WD_STYLE_TYPE.PARAGRAPH)
        evidence.base_style = document.styles["Normal"]
        evidence.font.name = "Aptos"
        evidence.font.size = Pt(9)
        evidence.font.bold = True
        evidence.font.color.rgb = RGBColor.from_string(NAVY)


def parse_markdown(markdown: str):
    title_match = re.search(r"^# (.+)$", markdown, re.M)
    title = title_match.group(1).strip() if title_match else "Institutional 29-Question Analysis"
    metadata = {}
    for key in ("Companies", "Evidence cut-off", "Method"):
        match = re.search(rf"^\*\*{re.escape(key)}:\*\*\s*(.+)$", markdown, re.M)
        if match:
            metadata[key] = match.group(1).strip()
    executive_match = re.search(r"## Executive assessment\s+([\s\S]*?)(?=\n---\s*\n)", markdown)
    executive = executive_match.group(1).strip() if executive_match else ""
    questions = []
    pattern = re.compile(r"^## (\d+)\.\s*(.+?)\s*$\n([\s\S]*?)(?=^## \d+\.|^---\s*$)", re.M)
    for match in pattern.finditer(markdown):
        questions.append((int(match.group(1)), match.group(2).strip(), match.group(3).strip()))
    conclusion_match = re.search(r"## Final comparative conclusion\s+([\s\S]*?)(?=\n---\s*\n)", markdown)
    conclusion = conclusion_match.group(1).strip() if conclusion_match else ""
    sources_match = re.search(r"## (?:Primary evidence used|Principal sources)\s+([\s\S]*)$", markdown)
    sources = sources_match.group(1).strip() if sources_match else ""
    return title, metadata, executive, questions, conclusion, sources


def render(input_path: Path, output_path: Path) -> None:
    markdown = input_path.read_text(encoding="utf-8")
    title, metadata, executive, questions, conclusion, sources = parse_markdown(markdown)
    if len(questions) != 29:
        raise ValueError(f"Expected 29 questions; found {len(questions)}")

    document = Document()
    configure_document(document)

    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title)
    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(metadata.get("Companies", ""))
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    meta_table = document.add_table(rows=0, cols=2)
    meta_table.autofit = False
    meta_table.columns[0].width = Inches(1.25)
    meta_table.columns[1].width = Inches(5.85)
    for key in ("Evidence cut-off", "Method"):
        if key not in metadata:
            continue
        cells = meta_table.add_row().cells
        cells[0].text = key
        cells[1].text = metadata[key]
        shade(cells[0], NAVY)
        shade(cells[1], LIGHT_GREY)
        for run in cells[0].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell)

    document.add_paragraph()
    document.add_heading("Executive assessment", level=1)
    for block in re.split(r"\n\s*\n", executive):
        add_rich_paragraph(document, block, shade_fill=LIGHT_BLUE)

    document.add_page_break()
    document.add_heading("Institutional due-diligence answers", level=1)
    for number, heading, body in questions:
        document.add_heading(f"{number}. {heading}", level=2)
        for block in re.split(r"\n\s*\n", body):
            if not block.strip():
                continue
            evidence = "Evidence state:" in block
            fill = None
            if evidence:
                fill = GREEN if "🟢" in block else AMBER if "🟠" in block else RED if "🔴" in block else LIGHT_GREY
            add_rich_paragraph(document, block.strip(), style="Evidence" if evidence else None, shade_fill=fill)

    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Final comparative conclusion", level=1)
    for block in re.split(r"\n\s*\n", conclusion):
        add_rich_paragraph(document, block, shade_fill=LIGHT_BLUE)

    document.add_heading("Primary evidence used", level=1)
    for line in sources.splitlines():
        line = line.strip()
        if line.startswith("-"):
            add_rich_paragraph(document, line[1:].strip(), style="List Bullet")
        elif line:
            add_rich_paragraph(document, line)

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "WEALTHOS  |  INSTITUTIONAL EQUITY RESEARCH"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(TEAL)
        footer = section.footer.paragraphs[0]
        footer.add_run("Evidence-bounded research  •  Page ")
        add_page_number(footer)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Usage: render_report_docx.py <report.md> [output.docx]")
    source = Path(sys.argv[1]).resolve()
    destination = Path(sys.argv[2]).resolve() if len(sys.argv) == 3 else source.with_suffix(".docx")
    render(source, destination)
    print(destination)
